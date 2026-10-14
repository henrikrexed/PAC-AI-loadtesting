#!/usr/bin/env python3
"""ISI-5720 — build the speaking-narrative Word doc from the committed markdown.

Extracts Part 3 (timed speaking narrative) of narrative-and-slide-copy.md and
renders it as a presenter-ready .docx: title page + per-slide headings (with
timings) + narration, with inline **bold** / *italic* / `code` preserved.

Deterministic, stdlib + python-docx only. Run from presentation-matrix/.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = Path("narrative-and-slide-copy.md")
OUT = Path("dist/PAC-there-is-no-spoon-narrative.docx")

MATRIX_GREEN = RGBColor(0x00, 0x8F, 0x3C)
INK = RGBColor(0x14, 0x14, 0x14)
MUTED = RGBColor(0x55, 0x55, 0x55)

# --- inline markdown -> runs -------------------------------------------------
TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


def add_runs(paragraph, text):
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.color.rgb = MATRIX_GREEN
        elif part.startswith("*") and part.endswith("*"):
            r = paragraph.add_run(part[1:-1]); r.italic = True
        else:
            paragraph.add_run(part)


def extract_part3(md):
    lines = md.splitlines()
    start = next(i for i, l in enumerate(lines) if l.startswith("# Part 3"))
    end = next(i for i, l in enumerate(lines) if i > start and l.startswith("# Part 4"))
    return lines[start + 1:end]


def main():
    md = SRC.read_text(encoding="utf-8")
    body = extract_part3(md)

    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11.5)

    # --- title page ---
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("THERE IS NO SPOON"); r.bold = True; r.font.size = Pt(30); r.font.color.rgb = INK
    st = doc.add_paragraph(); st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = st.add_run("Load-Testing the Real World"); r.italic = True; r.font.size = Pt(16); r.font.color.rgb = MATRIX_GREEN
    for line in (
        "Speaking narrative — presenter script",
        "Henrik Rexed · IsItObservable · PAC expert session (30 min)",
        "Theme: The Matrix · 21 slides + 3 discussion beats",
        "Source: presentation-matrix/narrative-and-slide-copy.md (Part 3)",
    ):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = p.add_run(line); rr.font.size = Pt(10.5); rr.font.color.rgb = MUTED
    doc.add_page_break()

    # --- narrative ---
    i = 0
    while i < len(body):
        line = body[i].rstrip()
        if not line or line == "---":
            i += 1
            continue
        if line.startswith("*") and line.endswith("*") and not line.startswith("###"):
            # the Part-3 intro italic note
            p = doc.add_paragraph(); add_runs(p, line)
            for rn in p.runs:
                rn.italic = True; rn.font.size = Pt(10.5); rn.font.color.rgb = MUTED
            i += 1
            continue
        if line.startswith("### "):
            head = line[4:].strip()
            h = doc.add_heading(level=2)
            is_disc = "💬" in head
            hr = h.add_run(head.replace("💬", "").strip())
            hr.font.color.rgb = MATRIX_GREEN if is_disc else INK
            # gather following non-blank, non-heading lines as the narration
            i += 1
            buf = []
            while i < len(body) and body[i].strip() and not body[i].startswith("### ") and body[i].strip() != "---":
                buf.append(body[i].strip())
                i += 1
            if buf:
                para = doc.add_paragraph()
                para.paragraph_format.space_after = Pt(10)
                add_runs(para, " ".join(buf))
            continue
        i += 1

    OUT.parent.mkdir(exist_ok=True)
    doc.save(OUT)
    # quick integrity report
    heads = sum(1 for l in body if l.startswith("### "))
    print(f"wrote {OUT} | slide/beat headings: {heads}")


if __name__ == "__main__":
    sys.exit(main())
